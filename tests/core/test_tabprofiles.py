import json
import math

import pandas as pd
import pytest

import pandaspro as cpd
from pandaspro.core.frame import FramePro, TabFrame

ROWS = [
    ('Staff HQ', 'IBRD/IDA', 'Open', 3),
    ('Staff HQ', 'IBRD/IDA', 'Term', 1),
    ('Staff HQ', 'IFC', 'Open', 1),
    ('Staff non-HQ', 'IBRD/IDA', 'Open', 2),
    ('Staff non-HQ', 'IFC', 'Term', 2),
    ('Staff non-HQ', 'MIGA', 'Open', 1),
]
ORG_ORDER = ['IBRD/IDA', 'IFC', 'MIGA']
SPEC = {
    'name': 'mydefault',
    'total_position': {'rows': 'last', 'cols': 'first'},
    'formats': {'count': '{:,.0f}', 'zero': '-', 'share': '{:.1%}', 'mean': '{:,.1f}'},
    'fields': {
        'org': {'order': ORG_ORDER, 'total_label': 'WBG', 'label': 'Organization', 'total_position': 'first'},
        'open_term': {'order': ['Open', 'Term'], 'share': {'value': 'Open', 'label': '% Open'}},
        'loc_new': {'order': ['Staff HQ', 'Staff non-HQ'], 'total_label': 'Total', 'label': 'Locations'},
        'grade': {'order': ['GA', 'GB'], 'total_label': 'All grades'},      # not used by the calls below
    },
}
EXPECTED_COLUMNS = [(org, inner) for org in ['WBG'] + ORG_ORDER for inner in ['Open', 'Term', '% Open']]


@pytest.fixture(autouse=True)
def clean_registries():
    def reset():
        cpd.clear_field_rules()
        cpd.clear_tab_defaults()
        for name in list(cpd.tab_profiles()):
            cpd.unregister_tab_profile(name)
    reset()
    yield
    reset()


def _people():
    records, upi = [], 0
    for loc, org, line, n in ROWS:
        for _ in range(n):
            upi += 1
            records.append({'upi': upi, 'loc_new': loc, 'org': org, 'open_term': line, 'salary': 100 + upi * 10})
    return FramePro(records)


def _same(one, two):
    pd.testing.assert_frame_equal(pd.DataFrame(one), pd.DataFrame(two))


# ---- registry

def test_register_list_get_unregister():
    cpd.register_tab_profile('mydefault', SPEC)
    assert list(cpd.tab_profiles()) == ['mydefault']
    assert cpd.tab_profile('mydefault')['fields']['org']['total_label'] == 'WBG'
    cpd.tab_profile('mydefault')['fields']['org']['total_label'] = 'tampered'    # a copy
    assert cpd.tab_profile('mydefault')['fields']['org']['total_label'] == 'WBG'
    cpd.register_tab_profile('mydefault', {'fields': {}})                        # replaces
    assert cpd.tab_profile('mydefault')['fields'] == {}
    cpd.unregister_tab_profile('mydefault')
    assert cpd.tab_profiles() == {}
    with pytest.raises(KeyError, match='is not registered'):
        cpd.tab_profile('mydefault')


@pytest.mark.parametrize('name', ['nototal', 'tsort_x', 'tdiff', 'display', 'numbers', 'style', 'copy',
                                  'cpdtab2_x', 'rename_agg'])
def test_profile_name_clashing_with_a_table_attribute_is_refused(name):
    with pytest.raises(ValueError, match='clashes with an existing table attribute'):
        cpd.register_tab_profile(name, SPEC)


@pytest.mark.parametrize('name', ['MyDefault', '1st', 'my-default', '', 'has space'])
def test_profile_name_must_be_lower_snake_case(name):
    with pytest.raises(ValueError, match='lower_snake_case'):
        cpd.register_tab_profile(name, SPEC)


def test_unknown_keys_warn_and_bad_values_raise():
    with pytest.warns(UserWarning, match="unknown keys ignored: \\['colour'\\]"):
        cpd.register_tab_profile('odd', {**SPEC, 'colour': 'red'})
    assert 'colour' not in cpd.tab_profile('odd')
    with pytest.warns(UserWarning, match="field 'org'"):
        cpd.register_tab_profile('odd', {'fields': {'org': {'ordre': []}}})
    with pytest.raises(ValueError, match='is not supported'):
        cpd.register_tab_profile('bad', {'formats': {'count': '%d'}})
    with pytest.raises(ValueError, match="'first' or 'last'"):
        cpd.register_tab_profile('bad', {'total_position': {'cols': 'left'}})
    with pytest.raises(ValueError, match="'first' or 'last'"):
        cpd.register_tab_profile('bad', {'fields': {'org': {'total_position': 'left'}}})


def test_load_profiles_from_a_file_and_a_folder(tmp_path):
    single = tmp_path / 'single.json'
    single.write_text(json.dumps(SPEC))
    assert cpd.load_tab_profiles(single) == ['mydefault']

    folder = tmp_path / 'profiles'
    folder.mkdir()
    (folder / 'by_file_name.json').write_text(json.dumps({'fields': {'org': {'total_label': 'All'}}}))
    (folder / 'several.json').write_text(json.dumps([{**SPEC, 'name': 'first_one'}, {**SPEC, 'name': 'second_one'}]))
    (folder / 'notes.txt').write_text('ignored')
    assert cpd.load_tab_profiles(folder) == ['by_file_name', 'first_one', 'second_one']
    assert set(cpd.tab_profiles()) == {'mydefault', 'by_file_name', 'first_one', 'second_one'}
    with pytest.raises(FileNotFoundError):
        cpd.load_tab_profiles(tmp_path / 'missing.json')


# ---- acceptance

def test_acceptance_chain_and_profile_argument():
    df = _people()
    cpd.register_tab_profile('mydefault', SPEC)
    plain = df.cpdtab2_loc_new___org__open_term
    assert type(plain) is FramePro and plain.columns[-1] == ('Total', '')        # plain stays plain
    assert 'WBG' not in plain.columns.get_level_values(0)

    mine = plain.mydefault
    assert isinstance(mine, TabFrame)
    assert list(mine.columns) == EXPECTED_COLUMNS
    assert list(mine.index) == ['Staff HQ', 'Staff non-HQ', 'Total']
    assert mine.index.name == 'Locations' and mine.columns.names[0] == 'Organization'
    assert mine.numbers.loc['Total', ('WBG', '% Open')] == 0.7
    assert type(mine.numbers) is pd.DataFrame
    assert '70.0%' in repr(mine) and '70.0%' in mine._repr_html_()

    _same(df.cpdtab2(index=['loc_new'], columns=['org', 'open_term'], profile='mydefault'), mine)
    renamed = df.cpdtab2('loc_new', ['org', 'open_term'], profile='mydefault', total_labels={'org': 'All'})
    assert renamed.columns[0] == ('All', 'Open')


def test_profile_on_an_aggregation_shortcut_uses_the_mean_format():
    cpd.register_tab_profile('mydefault', SPEC)
    df = _people()
    table = df.cpdtab2mean_loc_new__org__salary.mydefault
    assert list(table.columns) == ['WBG'] + ORG_ORDER
    assert table.numbers.loc['Total', 'WBG'] == pytest.approx(df['salary'].mean())
    assert table.display().loc['Total', 'WBG'] == f"{df['salary'].mean():,.1f}"
    assert table.display().loc['Staff HQ', 'MIGA'] == '-'                         # no rows: NaN shown as zero text


def test_profile_needs_a_cpdtab2_result():
    cpd.register_tab_profile('mydefault', SPEC)
    df = _people()
    assert not hasattr(df, 'mydefault')
    with pytest.raises(AttributeError, match='is a table profile'):
        df.mydefault
    with pytest.raises(AttributeError):
        df.cpdtab2s_loc_new__org.mydefault                                        # subtotal shortcut: not rebuildable
    with pytest.raises(KeyError, match='is not registered'):
        df.cpdtab2('loc_new', 'org', profile='nope')


# ---- precedence

def test_precedence_argument_then_field_rule_then_profile_setting():
    cpd.register_tab_profile('mine', {
        'total_position': 'last',
        'fields': {'org': {'total_position': 'first', 'total_label': 'WBG'}},
    })
    df = _people()
    by_rule = df.cpdtab2('loc_new', 'org', profile='mine')
    assert by_rule.columns[0] == 'WBG' and by_rule.index[-1] == 'Total'           # field rule beats profile setting
    by_arg = df.cpdtab2('loc_new', 'org', profile='mine', total_position={'cols': 'last'})
    assert by_arg.columns[-1] == 'WBG'                                            # argument beats field rule
    on_rows = df.cpdtab2('org', 'loc_new', profile='mine')
    assert on_rows.index[0] == 'WBG' and on_rows.columns[-1] == 'Total'           # follows the field's axis
    # several fields on one axis: the outermost field that declares a position decides (open_term has none)
    nested = df.cpdtab2('loc_new', ['open_term', 'org'], profile='mine')
    assert nested.columns[0] == ('Total', 'WBG')
    cpd.register_tab_profile('mine', {
        'fields': {'org': {'total_position': 'first'}, 'open_term': {'total_position': 'last'}},
    })
    assert df.cpdtab2('loc_new', ['open_term', 'org'], profile='mine').columns[-1] == ('Total', 'Total')
    assert df.cpdtab2('loc_new', ['org', 'open_term'], profile='mine').columns[0] == ('Total', 'Total')


def test_profile_ignores_global_rules_and_unused_field_rules():
    cpd.set_field_rules('org', total_label='GLOBAL', order=['MIGA'])
    cpd.set_tab_defaults(total_position='first')
    cpd.register_tab_profile('mine', {'fields': {'grade': {'total_label': 'All grades'}}})
    table = _people().cpdtab2('loc_new', 'org', profile='mine')
    assert list(table.columns) == ['IBRD/IDA', 'IFC', 'MIGA', 'Total']            # built-in defaults only
    assert list(table.index) == ['Staff HQ', 'Staff non-HQ', 'Total']
    assert _people().cpdtab2('loc_new', 'org').columns[0] == 'GLOBAL'             # globals still work without profile


def test_chaining_a_second_profile_replaces_the_first():
    cpd.register_tab_profile('mydefault', SPEC)
    cpd.register_tab_profile('other', {'fields': {'org': {'total_label': 'Bank Group'}}})
    df = _people()
    plain = df.cpdtab2_loc_new___org__open_term
    second = plain.mydefault.other
    _same(second, plain.other)
    _same(second, df.cpdtab2('loc_new', ['org', 'open_term'], profile='other'))
    assert second.columns[-1] == ('Bank Group', 'Total') and '% Open' not in second.columns.get_level_values(1)
    _same(second.mydefault, plain.mydefault)
    # explicit arguments of the original call survive a profile switch
    custom = df.cpdtab2('loc_new', 'org', totals='rows')
    assert 'WBG' not in custom.mydefault.columns and 'Total' in custom.mydefault.index


def test_profile_can_be_applied_after_other_operations():
    cpd.register_tab_profile('mydefault', SPEC)
    plain = _people().cpdtab2_loc_new___org__open_term
    _same(plain.copy().mydefault, plain.mydefault)
    _same(plain.mydefault.nototal.mydefault, plain.mydefault)


# ---- formats

def test_formatted_repr_numbers_and_display():
    cpd.register_tab_profile('mydefault', SPEC)
    table = _people().cpdtab2_loc_new___org__open_term.mydefault
    shown = table.display()
    assert shown.loc['Staff HQ'].tolist() == ['4', '1', '80.0%', '3', '1', '75.0%', '1', '-', '100.0%', '-', '-', '-']
    assert shown.loc['Staff non-HQ', ('IFC', '% Open')] == '0.0%'                 # a real 0% is not the zero text
    assert table.loc['Staff HQ', ('WBG', 'Open')] == 4 and str(table[('WBG', 'Open')].dtype) == 'int64'
    assert math.isnan(table.numbers.loc['Staff HQ', ('MIGA', '% Open')])

    cpd.register_tab_profile('two_decimals', {'formats': {'count': '{:,.0f}', 'share': '{:.2%}'},
                                              'fields': SPEC['fields']})
    other = _people().cpdtab2_loc_new___org__open_term.two_decimals
    assert other.display().loc['Total', ('IBRD/IDA', '% Open')] == '83.33%'
    assert other.display().loc['Staff HQ', ('IFC', 'Term')] == '0'                # no zero text declared

    big = FramePro({'a': ['x'] * 1234, 'b': ['y'] * 1234})
    assert '1,234' in repr(big.cpdtab2('a', 'b', profile='mydefault'))


def test_profile_without_formats_keeps_the_numeric_repr():
    cpd.register_tab_profile('bare', {'fields': SPEC['fields']})
    table = _people().cpdtab2_loc_new___org__open_term.bare
    assert '0.7' in repr(table) and '70.0%' not in repr(table)
    assert table.display().loc['Total', ('WBG', '% Open')] == '70.0%'            # display() still formats


def test_to_excel_number_formats(tmp_path):
    from openpyxl import load_workbook
    cpd.register_tab_profile('mydefault', SPEC)
    table = _people().cpdtab2_loc_new___org__open_term.mydefault
    table.to_excel(tmp_path / 'tab.xlsx')
    sheet = load_workbook(tmp_path / 'tab.xlsx').active
    assert sheet['B1'].value == 'WBG' and sheet['A2'].value == 'Locations'
    assert sheet['B3'].value == 4 and sheet['B3'].number_format == '#,##0;-#,##0;"-"'
    assert sheet['D3'].value == 0.8 and sheet['D3'].number_format == '0.0%'
    assert sheet['I3'].value == 0                                                 # numeric zero, shown as "-" by Excel

    mean = _people().cpdtab2mean_loc_new__org__salary.mydefault
    mean.to_excel(tmp_path / 'mean.xlsx')
    assert load_workbook(tmp_path / 'mean.xlsx').active['B2'].number_format == '#,##0.0;-#,##0.0;"-"'


# ---- plain stays plain

def test_legacy_outputs_are_unchanged_with_profiles_registered():
    df = _people()
    names = ['cpdtab2_loc_new__org', 'cpdtab2_loc_new___org__open_term', 'cpdtab2s_org__loc_new___open_term',
             'cpdtab2sum_loc_new__org__salary', 'cpdtab2mean_loc_new___org__salary', 'cpdtab2pctrow_loc_new__org']
    before = [getattr(df, name) for name in names]
    cpd.register_tab_profile('mydefault', SPEC)
    for name, old in zip(names, before):
        new = getattr(df, name)
        assert type(new) is FramePro and type(old) is FramePro
        _same(old, new)
        assert repr(old) == repr(new)
    # and the 1.2.0 method without a profile is not affected either
    table = df.cpdtab2('loc_new', 'org')
    assert list(table.columns) == ['IBRD/IDA', 'IFC', 'MIGA', 'Total'] and '0.' not in repr(table)

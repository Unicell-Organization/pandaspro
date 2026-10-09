"""Presentation of a cpdtab2 result: formatted strings, Styler, formatted Excel sheet.

The numbers in the table are never changed here; only how they are shown.
"""

from __future__ import annotations

import pandas as pd

from pandaspro.core.tools.tab2 import MARGIN_NAMES, TAB_META_KEY

COUNT_NUMBER_FORMAT = '#,##0;-#,##0;"-"'
DECIMAL_NUMBER_FORMAT = '#,##0.00;-#,##0.00;"-"'
PCT_NUMBER_FORMAT = '0.0%'


def _parts(label) -> tuple:
    return label if isinstance(label, tuple) else (label,)


def _label_sets(table: pd.DataFrame) -> tuple[set, set]:
    meta = table.attrs.get(TAB_META_KEY) or {}
    totals = set(MARGIN_NAMES) | set(meta.get('total_labels', {}).values())
    return totals, set(meta.get('pct_labels', []))


def _flags(labels, wanted: set) -> list[bool]:
    return [any(part in wanted for part in _parts(label)) for label in labels]


def format_number(value) -> str:
    if pd.isna(value) or value == 0:
        return '-'
    if float(value).is_integer():
        return f'{int(value):,}'
    return f'{value:,.2f}'


def format_pct(value) -> str:
    return '-' if pd.isna(value) else f'{value:.1%}'


def format_as(value, code: str | None) -> str:
    """Text for one cell. code: None (auto number), 'int', 'pctN' or 'numN' (N decimals)."""
    if code is None:
        return format_number(value)
    if pd.isna(value):
        return '-'
    if code == 'int':
        return '-' if round(value) == 0 else f'{round(value):,}'
    decimals = int(code[3:])
    return f'{value:.{decimals}%}' if code.startswith('pct') else f'{value:,.{decimals}f}'


def excel_format(value, code: str | None) -> str:
    if code is None:
        whole = value is None or float(value).is_integer()
        return COUNT_NUMBER_FORMAT if whole else DECIMAL_NUMBER_FORMAT
    if code == 'int':
        return COUNT_NUMBER_FORMAT
    decimals = int(code[3:])
    zeros = ('.' + '0' * decimals) if decimals else ''
    return f'0{zeros}%' if code.startswith('pct') else f'#,##0{zeros}'


def _axis_codes(labels, by_level: dict, pct_labels: set) -> list:
    """Format code per row / column: an explicit hint, else 'pct1' for share labels, else None."""
    codes = []
    for label in labels:
        parts, code = _parts(label), None
        for level, by_value in by_level.items():
            if int(level) < len(parts) and parts[int(level)] in by_value:
                code = by_value[parts[int(level)]]
        if code is None and any(part in pct_labels for part in parts):
            code = 'pct1'
        codes.append(code)
    return codes


def cell_codes(table: pd.DataFrame) -> tuple[list, list, str | None]:
    """(row codes, column codes, table default). A row code wins over a column code."""
    meta = table.attrs.get(TAB_META_KEY) or {}
    hints = meta.get('formats') or {}
    pct_labels = set(meta.get('pct_labels', []))
    return (_axis_codes(table.index, hints.get('rows', {}), pct_labels),
            _axis_codes(table.columns, hints.get('cols', {}), pct_labels),
            hints.get('default'))


def tab_display(table: pd.DataFrame) -> pd.DataFrame:
    """Same shape as the table, cells as formatted strings."""
    row_codes, col_codes, default = cell_codes(table)
    plain = pd.DataFrame(table)
    out = pd.DataFrame(index=plain.index, columns=plain.columns, dtype=object)
    for j, col_code in enumerate(col_codes):
        out.iloc[:, j] = [
            format_as(value, row_code or col_code or default)
            for value, row_code in zip(plain.iloc[:, j], row_codes)
        ]
    return out


def tab_style(table: pd.DataFrame):
    """Styler over the numbers: formatted, right-aligned, total rows / columns in bold."""
    totals, _ = _label_sets(table)
    plain = pd.DataFrame(table)
    total_rows, total_cols = _flags(plain.index, totals), _flags(plain.columns, totals)
    row_codes, col_codes, default = cell_codes(table)

    def bold(_):
        return pd.DataFrame(
            [['font-weight: bold' if (in_row or in_col) else '' for in_col in total_cols] for in_row in total_rows],
            index=plain.index, columns=plain.columns,
        )

    def formatter(code):
        return lambda value: format_as(value, code)

    styler = plain.style.format({label: formatter(code or default) for label, code in zip(plain.columns, col_codes)})
    for code in {code for code in row_codes if code}:
        rows = [label for label, row_code in zip(plain.index, row_codes) if row_code == code]
        styler = styler.format(formatter(code), subset=pd.IndexSlice[rows, :])
    return styler.set_properties(**{'text-align': 'right'}).apply(bold, axis=None)


def tab_to_excel(table: pd.DataFrame, path, sheet_name: str = 'Sheet1') -> None:
    """Write the table to a new workbook with openpyxl: numbers stay numbers, formats applied."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    totals, _ = _label_sets(table)
    plain = pd.DataFrame(table)
    n_header, n_index = plain.columns.nlevels, plain.index.nlevels
    total_rows, total_cols = _flags(plain.index, totals), _flags(plain.columns, totals)
    row_codes, col_codes, default = cell_codes(table)
    bold, center, right = Font(bold=True), Alignment(horizontal='center'), Alignment(horizontal='right')

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name

    # header: one row per column level, runs of the same outer label merged
    column_labels = [_parts(label) for label in plain.columns]
    for level in range(n_header):
        start = 0
        while start < len(column_labels):
            end = start
            while (level < n_header - 1 and end + 1 < len(column_labels)
                   and column_labels[end + 1][:level + 1] == column_labels[start][:level + 1]):
                end += 1
            header = sheet.cell(row=level + 1, column=n_index + start + 1, value=column_labels[start][level])
            header.font, header.alignment = bold, center
            if end > start:
                sheet.merge_cells(start_row=level + 1, start_column=n_index + start + 1,
                                  end_row=level + 1, end_column=n_index + end + 1)
            start = end + 1
    for position, name in enumerate(plain.index.names):
        sheet.cell(row=n_header, column=position + 1, value=name).font = bold

    # body
    for i, row_label in enumerate(plain.index):
        excel_row = n_header + 1 + i
        for position, part in enumerate(_parts(row_label)):
            label_cell = sheet.cell(row=excel_row, column=position + 1, value=part)
            if total_rows[i]:
                label_cell.font = bold
        for j in range(plain.shape[1]):
            value = plain.iat[i, j]
            if pd.isna(value):
                value = None
            elif hasattr(value, 'item'):
                value = value.item()
            data_cell = sheet.cell(row=excel_row, column=n_index + j + 1, value=value)
            data_cell.number_format = excel_format(value, row_codes[i] or col_codes[j] or default)
            data_cell.alignment = right
            if total_rows[i] or total_cols[j]:
                data_cell.font = bold

    workbook.save(path)
